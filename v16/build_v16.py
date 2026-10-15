"""Add operator work status to the validated v15 snapshot without changing raw sheets."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sqlite3
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from model_v16 import field, make_worklist

ROOT = Path(__file__).resolve().parent
SOURCE_NAMES = ('statistics.html', 'progress.html', 'case.html', 'work.html', 'app.js', 'app.css')
NAVY, BLUE, TEAL, PALE, WHITE = '102A43', '2864A8', '087F8C', 'EAF2F7', 'FFFFFF'
REQUIREMENTS = [
    ('REQ-01', '누적 연도별 신청·승인', '부분', '원장행·승인 상태 제공; 고유 사업/회차는 미확정'),
    ('REQ-02', '월별 신청·보고 통계', '구현', '관리 Tool/ORS 출처 분리 월별 통계'),
    ('REQ-03', '사업별 단계·종결', '부분', '2026 관리 53행·ORS 활동 169행; 회차 연결 미확정'),
    ('REQ-04', '보완 1·2·3차 이력', '부분', '원본 상태의 차수 표기·타임라인 제공; 미기록 가능'),
    ('REQ-05', '검토기관·담당자·3개 기관', '부분', '기관 역할 분리; 컨소시엄 3기관 자료 미수령'),
    ('REQ-06', '장기 미처리 경고', '부분', '30일 내부 점검; 법정·계약 기한 기준 미설정'),
    ('REQ-07', '월별 ORS Excel 갱신', '구현', 'ZIP 검증·원본 보존·자동 재생성'),
    ('REQ-08', '갱신 신청회차 보존', '부분', '수신본 이력 보존; 회차 고유 ID 미확정'),
    ('REQ-09', '최초 검토기관 보존', '부분', '관리 Tool 원본 보존; ORS 관장기관을 최초 검토기관으로 단정하지 않음'),
    ('REQ-10', '현장 방문·담당자', '부분', '원본 필드 조회; 별도 기관 입력자료 필요'),
    ('REQ-11', '기존 Tool 승계·디자인', '구현', '원본 74시트 유지, 작업현황·상세·그래프 추가'),
    ('REQ-12', 'Excel 우선·보안', '운영결정 변경', 'Excel 제공; 웹 공개는 이후 사용자 지시로 전환'),
]


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + '.building')
    temporary.write_text(json.dumps(value, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    os.replace(temporary, path)


def manual_notes(db):
    con = sqlite3.connect(db)
    con.row_factory = sqlite3.Row
    try:
        rows = con.execute('SELECT case_key,intake_status,opinion_status,opinion_text,guidance_text,guidance_status,next_action,owner,updated_at FROM workflow_notes').fetchall()
        return {row['case_key']: dict(row) for row in rows}
    finally:
        con.close()


def header(ws, title, subtitle):
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:P2')
    ws['A1'] = title
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A1'].font = Font(name='Malgun Gothic', size=21, bold=True, color=WHITE)
    ws['A1'].alignment = Alignment(vertical='center')
    ws.merge_cells('A3:P3')
    ws['A3'] = subtitle
    ws['A3'].font = Font(name='Malgun Gothic', size=10, color='50667A')
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[3].height = 28
    for letter in 'ABCDEFGHIJKLMNOP':
        ws.column_dimensions[letter].width = 13


def card(ws, first, label, number, note, color):
    ws.merge_cells(start_row=5, start_column=first, end_row=5, end_column=first + 2)
    ws.merge_cells(start_row=6, start_column=first, end_row=7, end_column=first + 2)
    ws.merge_cells(start_row=8, start_column=first, end_row=8, end_column=first + 2)
    for row in ws.iter_rows(min_row=5, max_row=8, min_col=first, max_col=first + 2):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=color)
    label_cell, number_cell, note_cell = (ws.cell(row, first) for row in (5, 6, 8))
    label_cell.value = label
    number_cell.value = number
    note_cell.value = note
    label_cell.font = Font(name='Malgun Gothic', size=10, bold=True, color='DDECF5')
    number_cell.font = Font(name='Malgun Gothic', size=24, bold=True, color=WHITE)
    note_cell.font = Font(name='Malgun Gothic', size=9, color='DDECF5')


def chart(ws, min_row, max_row, anchor, title, color):
    graph = BarChart()
    graph.title = title
    graph.style = 10
    graph.height = 8
    graph.width = 16
    graph.y_axis.title = '출처별 원본행'
    graph.add_data(Reference(ws, min_col=2, min_row=min_row, max_row=max_row), titles_from_data=True)
    graph.set_categories(Reference(ws, min_col=1, min_row=min_row + 1, max_row=max_row))
    graph.legend = None
    if graph.series:
        graph.series[0].graphicalProperties.solidFill = color
    ws.add_chart(graph, anchor)


def safe(value):
    if isinstance(value, str):
        if value[:1] in ('=', '+', '@'):
            value = "'" + value
        return value[:32767]
    return value


def workbook(source, output, work):
    wb = load_workbook(source)
    for name in ('업무 운영 대시보드', '담당자 작업현황', '요구사항 점검'):
        if name in wb.sheetnames:
            raise ValueError(f'v16 시트가 이미 있습니다: {name}')
    dashboard = wb.create_sheet('업무 운영 대시보드', 2)
    queue = wb.create_sheet('담당자 작업현황', 3)
    requirements = wb.create_sheet('요구사항 점검', 4)
    summary = work['summary']
    header(dashboard, '업무 운영 | 단계·다음 조치',
           '관리 Tool과 ORS의 출처별 원본행 대기열입니다. 두 출처의 합은 고유 신청 수가 아닙니다.')
    card(dashboard, 1, '관리 Tool 2026', summary['management_rows'], '수기 의견·안내 입력 대상', NAVY)
    card(dashboard, 5, 'ORS 2026 활동', summary['ors_activity_rows'], '최신 산업부문 원본행', BLUE)
    card(dashboard, 9, '내부 30일 점검', summary['internal_30day_check'], '법정 기한 초과 판정 아님', TEAL)
    card(dashboard, 13, '담당자 미기록', summary['owner_missing'], '기관명과 담당자는 구분', NAVY)
    dashboard['A11'] = '담당자 입력 다음 조치'
    dashboard['B11'] = summary['manual_actions']
    dashboard['D11'] = '의견 내용 입력'
    dashboard['E11'] = summary['manual_opinions']
    dashboard['G11'] = '사무국 안내 입력'
    dashboard['H11'] = summary['manual_guidance']
    dashboard['J11'] = '보완 단계 원본행'
    dashboard['K11'] = summary['supplement']
    dashboard['A13'], dashboard['B13'] = '단계', '원본행'
    for index, (label, count) in enumerate(summary['stages'].items(), 14):
        dashboard.cell(index, 1, label)
        dashboard.cell(index, 2, count)
    stage_end = 13 + len(summary['stages'])
    dashboard['A27'], dashboard['B27'] = '확인 우선순위', '원본행'
    for index, (label, count) in enumerate(summary['priority'].items(), 28):
        dashboard.cell(index, 1, label)
        dashboard.cell(index, 2, count)
    chart(dashboard, 13, stage_end, 'D14', '2026 출처별 단계 현황', BLUE)
    chart(dashboard, 27, 27 + len(summary['priority']), 'J14', '담당자 확인 우선순위', TEAL)
    dashboard.merge_cells('A34:P36')
    dashboard['A34'] = ('다음 조치의 출처를 반드시 확인하세요. “담당자 입력”만 실제 업무 입력이고 “원본 기반 내부 제안”은 확인을 돕는 제안입니다. '
                        '법정·계약 기한은 미설정이며, 30일은 내부 점검 기준입니다. 담당자 입력은 2026 접수·검토·안내 시트에서 편집합니다.')
    dashboard['A34'].alignment = Alignment(wrap_text=True, vertical='center')
    dashboard['A34'].fill = PatternFill('solid', fgColor=PALE)
    dashboard['A34'].font = Font(name='Malgun Gothic', size=10, color=NAVY)
    dashboard['A39'], dashboard['B39'] = '관리 Tool 검토기관', '업무행'
    for index, (label, count) in enumerate(summary['management_review_org'].items(), 40):
        dashboard.cell(index, 1, label)
        dashboard.cell(index, 2, count)
    chart(dashboard, 39, 39 + len(summary['management_review_org']), 'E40', '관리 Tool 검토기관별 업무행', BLUE)
    dashboard.freeze_panes = 'A4'
    dashboard.sheet_view.zoomScale = 85

    headers = ['확인 우선순위','출처','현재 단계','사업·신청명','신청자','접수일','최근 이력일','경과일','보완 차수','업무 담당자','담당 근거',
               '검토기관','관장기관','다음 조치','조치 출처','특이사항·근거','의견 상태','안내 상태','현장 방문','내부 30일 점검','법정·계약 기한','원본/업무입력','업무 연결키']
    queue.append(headers)
    queue.sheet_view.showGridLines = False
    queue.freeze_panes = 'D2'
    queue.row_dimensions[1].height = 34
    for cell in queue[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Malgun Gothic', size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True)
    input_sheet = wb['2026 접수·검토·안내']
    input_rows = {str(input_sheet.cell(n, 1).value): n for n in range(6, input_sheet.max_row + 1)
                  if input_sheet.cell(n, 1).value}
    for row in work['rows']:
        queue.append([safe(x) for x in [row['priority'], row['source'], row['stage'], row['title'], row['applicant'],
                     row['intake_date'], row['last_date'], row['age_days'], row['supplement_round'], row['owner'],
                     row['owner_origin'], row['review_org'], row['authority_org'], row['next_action'],
                     row['action_origin'], f"{row['special_origin']} · {row['special']}",
                     row['opinion_status'], row['guidance_status'], row['site_visit'],
                     '내부 점검' if row['internal_30day_check'] else '', row['due_basis'],
                     row['raw_excel_sheet'] or '2026 접수·검토·안내', row['id']]])
        excel_row = queue.max_row
        if row['source'] == '관리 Tool':
            target = f"#'2026 접수·검토·안내'!A{input_rows[row['id']]}"
        else:
            target = f"#'{row['raw_excel_sheet']}'!A{int(row['source_row']) + 4}"
        queue.cell(excel_row, 22).hyperlink = target
        queue.cell(excel_row, 22).font = Font(name='Malgun Gothic', color=BLUE, underline='single')
        queue.cell(excel_row, 1).fill = PatternFill('solid', fgColor='FFF2D7' if row['priority'] != '일반 확인' else PALE)
    table = Table(displayName='OperatorWorklist2026', ref=f'A1:W{queue.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    queue.add_table(table)
    queue.auto_filter.ref = table.ref
    widths = [17,13,15,55,26,15,16,10,14,19,18,22,20,55,20,56,17,17,18,18,24,24,31]
    for i, width in enumerate(widths, 1):
        queue.column_dimensions[queue.cell(1, i).column_letter].width = width

    requirements.append(['회의 요구 ID','요구사항','v16 상태','검증 근거·남는 조건'])
    for row in REQUIREMENTS:
        requirements.append(row)
    requirements.sheet_view.showGridLines = False
    requirements.freeze_panes = 'A2'
    for cell in requirements[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Malgun Gothic', bold=True, color=WHITE)
    for column, width in zip('ABCD', [18,34,18,88]):
        requirements.column_dimensions[column].width = width
    requirements.auto_filter.ref = f'A1:D{requirements.max_row}'
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix('.building.xlsx')
    wb.save(temporary)
    wb.close()
    with zipfile.ZipFile(temporary) as archive:
        bad = archive.testzip()
        charts = sum(name.startswith('xl/charts/chart') and name.endswith('.xml') for name in archive.namelist())
    if bad or charts < 12:
        raise RuntimeError(f'Excel 무결성/차트 검사 실패: {bad}, {charts}')
    os.replace(temporary, output)
    return {'sheets': 77, 'charts': charts, 'work_rows': len(work['rows'])}


def build(v15, db, output):
    progress = read(v15 / 'data' / 'progress.json')
    management_details = read(v15 / 'details' / 'management.json')
    ors_details = {path.stem: read(path) for path in sorted((v15 / 'details').glob('ORS*.json'))}
    notes = manual_notes(db)
    work = make_worklist(progress, management_details, ors_details, notes)
    output.mkdir(parents=True, exist_ok=True)
    shutil.copytree(v15 / 'data', output / 'data', dirs_exist_ok=True)
    shutil.copytree(v15 / 'details', output / 'details', dirs_exist_ok=True)
    for group in ors_details.values():
        for detail in group.values():
            detail['first_agency'] = field(detail, '최초 검토기관')
            detail['reviewer'] = field(detail, '검토기관', '평가기관')
            detail['agency_role'] = '관장기관'
    for sheet, group in ors_details.items():
        write(output / 'details' / f'{sheet}.json', group)
    for name in SOURCE_NAMES:
        shutil.copy2(ROOT / name, output / name)
    write(output / 'data' / 'work.json', work)
    excel = workbook(v15 / '외부사업_통합관리_v15.xlsx',
                     output / '외부사업_통합관리_v16.xlsx', work)
    digest = hashlib.sha256()
    digest.update(read(v15 / 'manifest.json')['version'].encode())
    digest.update(json.dumps(work, ensure_ascii=False, sort_keys=True).encode('utf-8'))
    for name in ('model_v16.py', 'build_v16.py', *SOURCE_NAMES):
        digest.update((ROOT / name).read_bytes())
    manifest = {**read(v15 / 'manifest.json'), 'version': digest.hexdigest()[:16],
                'v15_version': read(v15 / 'manifest.json')['version'],
                'work_source_rows': len(work['rows']), 'workbook_sheets': excel['sheets']}
    write(output / 'manifest.json', manifest)
    return {'output': str(output), 'manifest': manifest, 'work': work['summary'], 'excel': excel}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--v15', type=Path, required=True)
    parser.add_argument('--db', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(build(args.v15, args.db, args.output), ensure_ascii=False))
