"""Verify the new worklist and that every v15 workbook source cell survived."""
from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def verify(site):
    old, root = site / 'v15', site / 'v16'
    progress = read(root / 'data' / 'progress.json')
    work = read(root / 'data' / 'work.json')
    manifest = read(root / 'manifest.json')
    assert manifest['work_source_rows'] == len(work['rows']) == 222
    assert work['summary']['management_rows'] == len(progress['management']) == 53
    assert work['summary']['ors_activity_rows'] == progress['overview']['ors_industry_2026_activity'] == 169
    assert work['summary']['source_rows'] == 222
    assert sum(work['summary']['priority'].values()) == 222
    assert sum(work['summary']['stages'].values()) == 222
    assert sum(work['summary']['management_review_org'].values()) == 53
    assert work['freshness'] == manifest['freshness']
    assert len({(row['source'], row['id']) for row in work['rows']}) == 222
    management = read(root / 'details' / 'management.json')
    ors = {path.stem: read(path) for path in (root / 'details').glob('ORS*.json')}
    assert len(ors) == 20
    assert sum(len(group) for group in ors.values()) == len(progress['ors']) == 5998
    for row in work['rows']:
        link = urlsplit(row['detail_url'])
        assert link.path == 'case.html'
        query = parse_qs(link.query)
        if row['source'] == '관리 Tool':
            assert management[query['id'][0]]['id'] == row['id']
        else:
            detail = ors[query['sheet'][0]][query['row'][0]]
            assert detail['id'] == row['id']
            raw = {x['label']: x['value'] for x in detail['raw_fields']}
            assert detail['first_agency'] == str(raw.get('최초 검토기관') or '').strip()
            assert detail['reviewer'] == str(raw.get('검토기관') or raw.get('평가기관') or '').strip()
        assert row['action_origin'] in ('담당자 입력', '원본 기반 내부 제안')
        assert row['due_basis'] == '법정·계약 기한 기준 미설정'
    for name in ('statistics.html', 'progress.html', 'work.html', 'case.html', 'app.js', 'app.css'):
        assert (root / name).stat().st_size > 100, name
    source_path = old / '외부사업_통합관리_v15.xlsx'
    target_path = root / '외부사업_통합관리_v16.xlsx'
    with zipfile.ZipFile(target_path) as archive:
        assert archive.testzip() is None
        charts = sum(x.startswith('xl/charts/chart') and x.endswith('.xml') for x in archive.namelist())
        assert charts >= 12
    source = load_workbook(source_path, read_only=True, data_only=False)
    target = load_workbook(target_path, read_only=True, data_only=False)
    try:
        assert len(source.sheetnames) == 74 and len(target.sheetnames) == 77
        assert set(source.sheetnames).issubset(target.sheetnames)
        assert target['담당자 작업현황'].max_row == 223
        assert target['담당자 작업현황'].max_column == 23
        assert target['요구사항 점검'].max_row == 13
        compared = 0
        for name in source.sheetnames:
            before, after = source[name], target[name]
            assert (before.max_row, before.max_column) == (after.max_row, after.max_column), name
            for left, right in zip(before.iter_rows(values_only=True), after.iter_rows(values_only=True)):
                assert left == right, f'원본 셀 불일치: {name}'
                compared += len(left)
    finally:
        source.close()
        target.close()
    return {'work_rows': len(work['rows']), 'management': 53, 'ors_activity': 169,
            'source_sheets_preserved': 74, 'source_cells_compared': compared,
            'new_sheets': 3, 'charts': charts, 'version': manifest['version']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--site', type=Path, required=True)
    print(json.dumps(verify(parser.parse_args().site), ensure_ascii=False))
