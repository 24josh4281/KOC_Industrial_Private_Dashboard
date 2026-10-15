"""Incrementally refresh the checked v13 Excel from a DB snapshot.

The large, source-complete workbook is the immutable starting point. New ORS
occurrences are appended; the 2026 workflow and activity views are refreshed.
This avoids re-exporting the full workbook through the artifact runtime.
"""
import argparse
import base64
import gzip
import json
import os
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / '외부사업_통합관리_전체자료_v13_검증완료.xlsx'
NAVY = '102A43'


def safe_cell(value):
    if isinstance(value, str) and value[:1] in ('=', '+', '@'):
        value = "'" + value
    if isinstance(value, str) and len(value) > 32767:
        return value[:32767]
    return value


def column_name(n):
    result = ''
    while n:
        n, rem = divmod(n - 1, 26)
        result = chr(65 + rem) + result
    return result


def append_ors(workbook, catalog):
    index = workbook['원본 안내']
    specials = workbook['원본 특수셀']
    baseline_count = len([name for name in workbook.sheetnames if name.startswith('ORS') and name[3:].isdigit()])
    if baseline_count > len(catalog):
        raise ValueError('DB ORS 원본 건수가 기준 Excel보다 적습니다. 원본 삭제 가능성을 확인하세요.')
    for n, item in enumerate(catalog[baseline_count:], start=baseline_count + 1):
        name = f'ORS{n:02d}'
        if name in workbook:
            raise ValueError(f'중복 Excel 시트: {name}')
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for row in (
            ('원본 ZIP', item['container'], None, None),
            ('원본 XLS', item['filename'], item['sheet'], None),
            ('읽는 법', 'A열은 원본 행번호 · B열부터 원본 A열~', 'XLS 저장값', '원본 XLS 수식·서식은 ZIP 참고'),
        ):
            sheet.append(row)
        sheet.append(('원본 행', '원본 A열'))
        for cell in sheet[4]:
            cell.fill = PatternFill('solid', fgColor=NAVY)
            cell.font = Font(name='Malgun Gothic', size=10, bold=True, color='FFFFFF')
        sheet.freeze_panes = 'A5'
        sheet.column_dimensions['A'].width = 12
        sheet.column_dimensions['B'].width = 30
        for source_row, cells in item['rows']:
            if source_row + 4 > 1048576:
                raise ValueError(f'{item["filename"]}: Excel 최대 행 수 초과')
            sheet.cell(source_row + 4, 1, source_row)
            for col, value, *_ in cells:
                if col + 1 > 16384:
                    raise ValueError(f'{item["filename"]}: Excel 최대 열 수 초과')
                if value == '':
                    specials.append((item['container'] + ' / ' + item['filename'], item['sheet'],
                                     f'{column_name(col)}{source_row}', 'ORS 빈 문자열 저장', None))
                elif isinstance(value, str) and ('\r' in value or len(value) > 32767):
                    specials.append((item['container'] + ' / ' + item['filename'], item['sheet'],
                                     f'{column_name(col)}{source_row}',
                                     'ORS 줄바꿈 원문(JSON)' if '\r' in value else 'ORS 장문 원문(JSON)',
                                     json.dumps(value, ensure_ascii=False)))
                if value not in (None, ''):
                    sheet.cell(source_row + 4, col + 1, safe_cell(value))
        index.append((name, item['container'], item['filename'] + ' / ' + item['sheet'], item['cells'], 0))
    index['A2'] = f'원본01~36은 관리 Excel, ORS01~{len(catalog):02d}은 ZIP 추출 XLS의 저장값입니다. 원본 서식·개체는 동봉 원본을 확인하세요.'
    return len(catalog) - baseline_count


def refresh_workflow(workbook, data):
    sheet = workbook['2026 접수·검토·안내']
    count = 0
    for row in range(6, sheet.max_row + 1):
        key = sheet.cell(row, 1).value
        if not key:
            continue
        note = data['workflow'].get(key)
        if note is None:
            raise ValueError(f'업무 연결키가 스냅샷에 없습니다: {key}')
        for col, field in enumerate(('intake_status', 'opinion_status', 'opinion_text', 'guidance_text',
                                     'guidance_status', 'next_action', 'owner', 'updated_at'), start=13):
            sheet.cell(row, col, safe_cell(note.get(field) or ''))
        count += 1
    if count != len(data['cases_2026']):
        raise ValueError(f'2026 원장 행 수 불일치: Excel {count} / DB {len(data["cases_2026"])}')
    return count


def refresh_activity(workbook, data):
    sheet = workbook['2026 ORS 이력후보']
    old_max = sheet.max_row
    if old_max >= 5:
        sheet.delete_rows(5, old_max - 4)
    latest = data.get('refresh_snapshot', {}).get('latest_ors_date', '')
    sheet['A1'] = f'2026년 ORS 이력 확인 후보 · 최신 수신 {latest}'
    activity = data.get('ors_2026_activity', [])
    for item in activity:
        sheet.append(tuple(safe_cell(value) for value in (
            item['source_file'], item['source_row'], item['kind'], item['number'], item['title'],
            item['status'], '\n'.join(item['events_2026']), item['event_count'], item['source_sha256'])))
    if 'OrsActivity2026' in sheet.tables:
        if activity:
            sheet.tables['OrsActivity2026'].ref = f'A4:I{len(activity) + 4}'
        else:
            del sheet.tables['OrsActivity2026']
    elif activity:
        table = Table(displayName='OrsActivity2026', ref=f'A4:I{len(activity) + 4}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    for row in range(5, len(activity) + 5):
        sheet.cell(row, 7).alignment = Alignment(wrap_text=True, vertical='top')
    return len(activity)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--data', type=Path, required=True)
    parser.add_argument('--ors', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if not BASE.exists():
        raise FileNotFoundError(f'검증된 기준 Excel이 없습니다: {BASE}')
    data = json.loads(args.data.read_text(encoding='utf-8'))
    catalog = json.loads(gzip.decompress(base64.b64decode(args.ors.read_text(encoding='ascii'))))
    workbook = load_workbook(BASE)
    try:
        new_ors = append_ors(workbook, catalog)
        cases = refresh_workflow(workbook, data)
        activity = refresh_activity(workbook, data)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        temp = args.output.with_suffix('.building.xlsx')
        workbook.save(temp)
        with zipfile.ZipFile(temp) as archive:
            bad = archive.testzip()
            charts = sum('/charts/chart' in name and name.endswith('.xml') for name in archive.namelist())
        if bad or charts < 4:
            temp.unlink(missing_ok=True)
            raise RuntimeError(f'Excel ZIP/차트 검사 실패: bad={bad}, charts={charts}')
        os.replace(temp, args.output)
        print(json.dumps({'output': str(args.output), 'ors_occurrences': len(catalog),
                          'new_ors_sheets': new_ors, 'cases_2026': cases,
                          'ors_activity_2026': activity, 'chart_files': charts}, ensure_ascii=False))
    finally:
        workbook.close()


if __name__ == '__main__':
    main()
