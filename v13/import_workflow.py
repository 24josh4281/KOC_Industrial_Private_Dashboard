"""Import edited 2026 review/guidance fields with an audit trail; originals stay immutable."""
import argparse
import csv
import hashlib
import json
import shutil
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FIELDS = ['intake_status','opinion_status','opinion_text','guidance_text','guidance_status','next_action','owner','updated_at']
DEFAULTS = {'intake_status':'확인 필요','opinion_status':'미기록','guidance_status':'미기록'}
ALLOWED = {
    'intake_status': {'확인 필요','접수 확인','인입 보류','해당 없음'},
    'opinion_status': {'미기록','검토 예정','검토 중','의견 정리','검토 완료','확인 필요'},
    'guidance_status': {'미기록','안내 초안','내부 확인','안내 완료','해당 없음'},
}
EXCEL_COLS = {'업무 연결키':'case_key','접수 확인':'intake_status','의견 검토 상태':'opinion_status','의견·검토 내용':'opinion_text','사무국 안내 내용':'guidance_text','안내 상태':'guidance_status','다음 조치':'next_action','업무 담당자':'owner','수정일':'updated_at'}


def read_records(path):
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
    if path.suffix.lower()!='.xlsx': raise ValueError('CSV 또는 XLSX만 지원합니다.')
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=False)
    if '2026 접수·검토·안내' not in wb: raise ValueError('2026 접수·검토·안내 시트가 없습니다.')
    ws=wb['2026 접수·검토·안내']
    headers=[EXCEL_COLS.get(c.value) for c in ws[5]]
    rows=[]
    for excel_row in ws.iter_rows(min_row=6):
        item={}
        for key,cell in zip(headers,excel_row):
            if not key: continue
            if cell.data_type in ('f','e'): raise ValueError(f'입력열의 수식/오류 금지: {cell.coordinate}')
            item[key]=cell.value.isoformat() if hasattr(cell.value,'isoformat') else cell.value
        if item.get('case_key'): rows.append(item)
    wb.close()
    return rows


def main():
    ap=argparse.ArgumentParser(description='2026 접수·의견·사무국 안내를 검증 후 DB 반영')
    ap.add_argument('file',type=Path)
    ap.add_argument('--commit',action='store_true')
    ap.add_argument('--db',type=Path,help='별도 시험 DB. 기본값은 state/working.sqlite3')
    a=ap.parse_args();path=a.file.resolve(strict=True)
    file_sha=hashlib.sha256(path.read_bytes()).hexdigest()
    records=read_records(path)
    dbpath=a.db
    if dbpath is None:
        state=ROOT/'state';state.mkdir(exist_ok=True)
        dbpath=state/'working.sqlite3'
        if not dbpath.exists():
            if a.commit:shutil.copy2(ROOT/'external_projects_v13.sqlite3',dbpath)
            else:dbpath=ROOT/'external_projects_v13.sqlite3'
    con=sqlite3.connect(dbpath)
    known={x[0] for x in con.execute("SELECT case_key FROM management_records WHERE apply_year='2026'")}
    if len(known)!=53: raise ValueError(f'2026 원장 범위가 바뀌었습니다: {len(known)}행. 새 판본 재생성 필요')
    seen=set();changes=[]
    for item in records:
        key=str(item['case_key']).strip()
        if key in seen: raise ValueError(f'업무 연결키 중복: {key}')
        seen.add(key)
        if key not in known: raise ValueError(f'원본 연결키를 찾을 수 없음: {key}')
        old=con.execute('SELECT '+','.join(FIELDS)+' FROM workflow_notes WHERE case_key=?',(key,)).fetchone()
        old=dict(zip(FIELDS,old)) if old else {f:DEFAULTS.get(f,'') for f in FIELDS}
        for field in FIELDS:
            new=str(item.get(field) or DEFAULTS.get(field,'')).strip()
            if len(new)>10000: raise ValueError(f'{key}: {field} 1만자 초과')
            if field in ALLOWED and new not in ALLOWED[field]: raise ValueError(f'{key}: {field} 허용되지 않은 값 {new}')
            if new!=str(old.get(field) or ''):
                changes.append({'case_key':key,'field':field,'before':old.get(field) or '', 'after':new})
    duplicate=con.execute('SELECT 1 FROM import_runs WHERE input_sha=?',(file_sha,)).fetchone() is not None
    report={'input':path.name,'sha256':file_sha,'rows':len(records),'changed_fields':len(changes),'changed_cases':len({x['case_key'] for x in changes}),'duplicate_input':duplicate,'mode':'commit' if a.commit else 'dry_run','changes':changes}
    reports=ROOT/'imports';reports.mkdir(exist_ok=True)
    name=reports/f'workflow_{file_sha[:12]}_{"commit" if a.commit else "preview"}.json'
    if a.commit and not duplicate and changes:
        at=datetime.now(timezone.utc).isoformat()
        backup_dir=(ROOT if a.db is None else dbpath.parent)/'backups';backup_dir.mkdir(exist_ok=True)
        backup=backup_dir/f'{datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")}_{file_sha[:12]}_before.sqlite3'
        backup_con=sqlite3.connect(backup);con.backup(backup_con);backup_con.close()
        report['backup']=str(backup)
        con.execute('BEGIN IMMEDIATE')
        for item in records:
            key=str(item['case_key']).strip()
            values=[str(item.get(f) or DEFAULTS.get(f,'')).strip() for f in FIELDS]
            if any(x['case_key']==key for x in changes):
                con.execute('INSERT INTO workflow_notes(case_key,'+','.join(FIELDS)+') VALUES('+','.join('?' for _ in range(len(FIELDS)+1))+') ON CONFLICT(case_key) DO UPDATE SET '+','.join(f+'=excluded.'+f for f in FIELDS),(key,*values))
        con.executemany('INSERT INTO workflow_history(case_key,changed_at,field,before_text,after_text,import_sha) VALUES(?,?,?,?,?,?)',[(x['case_key'],at,x['field'],x['before'],x['after'],file_sha) for x in changes])
        con.execute('INSERT INTO import_runs VALUES(?,?,?,?)',(file_sha,'workflow',at,json.dumps(report,ensure_ascii=False)))
        con.commit()
    con.close()
    if not name.exists(): name.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'report':str(name),'rows':len(records),'changed_fields':len(changes),'duplicate':duplicate,'mode':report['mode']},ensure_ascii=False))


if __name__=='__main__': main()
