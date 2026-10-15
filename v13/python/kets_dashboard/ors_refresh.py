"""Plan safe local ORS imports. No workbook writes and no network access."""
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import hashlib, io, json, re, zipfile

CONFIG=json.loads((Path(__file__).resolve().parents[2]/'config/ors_refresh.json').read_text(encoding='utf-8'))
SEOUL=timezone(timedelta(hours=9))

def normalized(value):
    if value is None:return ''
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value).strip()

def header(value):
    return re.sub(r'\s+','',normalized(value)).lower()

def dated_filename(name):
    matches=re.findall(r'(?<!\d)(20\d{6})(?:\d{6})?(?!\d)',name)
    dates=[]
    for value in matches:
        try:dates.append(datetime.strptime(value,'%Y%m%d').date().isoformat())
        except ValueError:pass
    return max(dates) if dates else None

def classify(name,headers):
    h={header(x) for x in headers};n=header(name)
    if header('방법론명') in h and header('신청일자') in h:return 'methodology'
    if header('사업명') in h and header('신청일') in h and header('진행상태') in h:
        if '국외' in n or '해외' in n:return None
        return 'disclosure_credit_domestic' if header('모니터링차수') in h else 'disclosure_project_domestic'
    region='foreign' if '국외' in n or '해외' in n else 'domestic' if '국내' in n else None
    if not region:return None
    if header('[모니터링정보]모니터링 차수') in h:return region+'_credit'
    if header('[신청정보] 신청사업명') in h and header('신청유형') in h:return region+'_project'
    return None

def read_exports(paths):
    """Read XLS/XLSX or contained files without extracting archive paths."""
    result=[]
    def read(payload,name,container=None):
        digest=hashlib.sha256(payload).hexdigest();tables=[]
        if name.lower().endswith('.xls'):
            try:import xlrd
            except ImportError as exc:raise RuntimeError('XLS 읽기에는 requirements-ingest.txt의 xlrd가 필요합니다.') from exc
            wb=xlrd.open_workbook(file_contents=payload)
            for sh in wb.sheets():
                rows=[]
                for r in range(sh.nrows):
                    vals=[]
                    for c in range(sh.ncols):
                        cell=sh.cell(r,c);v=cell.value
                        if cell.ctype==xlrd.XL_CELL_DATE:v=xlrd.xldate_as_datetime(v,wb.datemode).isoformat()
                        if cell.ctype==xlrd.XL_CELL_ERROR:raise ValueError(f'ORS 수신 파일 오류 셀: {name}/{sh.name}/{r+1}/{c+1}')
                        vals.append(v)
                    rows.append(vals)
                tables.append((sh.name,rows))
        else:
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=False)
            for sh in wb:
                rows=[]
                for row in sh:
                    if any(c.data_type in ('f','e') for c in row):raise ValueError('ORS 수신 파일은 수식·오류 없는 다운로드 값이어야 합니다.')
                    rows.append([c.value.isoformat() if isinstance(c.value,(date,datetime)) else c.value for c in row])
                tables.append((sh.title,rows))
            wb.close()
        for title,rows in tables:
            candidates=[(sum(header(c) in {header(x) for x in ['번호','신청유형','방법론명','사업명','처리상태','진행상태','관장기관']} for c in row),i) for i,row in enumerate(rows[:10])]
            score,hi=max(candidates,default=(0,0))
            if score<2:continue
            headers=rows[hi];kind=classify(name+' '+title,headers)
            result.append({'filename':name,'sheet':title,'container':container,'sha256':digest,'as_of':dated_filename(name),'kind':kind,
                'columns':[{'col':i+1,'label':normalized(v)} for i,v in enumerate(headers) if normalized(v)],
                'rows':[{'row':i+1,'cells':[{'col':j+1,'value':v} for j,v in enumerate(row) if v is not None and v!='']} for i,row in enumerate(rows) if i>hi and any(v is not None and v!='' for v in row)]})
    for path in sorted(map(Path,paths)):
        if path.name.startswith('~$'):continue
        if path.suffix.lower()=='.zip':
            with zipfile.ZipFile(path) as z:
                if sum(x.file_size for x in z.infolist())>200_000_000:raise ValueError('ORS ZIP 압축 해제 예상 용량이 200 MB를 초과합니다.')
                for item in z.infolist():
                    if not item.is_dir() and item.filename.lower().endswith(('.xls','.xlsx')):read(z.read(item),item.filename,path.name)
        elif path.suffix.lower() in ('.xls','.xlsx'):read(path.read_bytes(),path.name)
    return result

def column_map(columns):
    groups=defaultdict(list)
    for c in columns:
        if header(c['label']):groups[header(c['label'])].append(c['col'])
    return {k:v[0] for k,v in groups.items() if len(v)==1},[k for k,v in groups.items() if len(v)>1]

def identity(row,cols,keys):
    cells={c['col']:c for c in row['cells']}
    vals=[normalized(cells.get(cols.get(header(k)),{}).get('value')) for k in keys]
    return tuple(vals) if all(vals) else None

def make_plan(exports,sheets,baseline_by_book,config=CONFIG,as_of=None):
    as_of=as_of or datetime.now(SEOUL).date().isoformat()
    """Unique composite matching; preserve order/manual cells and reject stale inputs."""
    writes=[];events=[];pending=[];seen_files=set();claims=defaultdict(list)
    for source in exports:
        dedup=(source['sha256'],source['sheet'])
        if dedup in seen_files:
            events.append({'filename':source['filename'],'source_date':source['as_of'],'status':'duplicate_file','rows':len(source['rows']),'matched':0,'updated_cells':0});continue
        seen_files.add(dedup)
        spec=config['datasets'].get(source['kind']);sc,sd=column_map(source['columns'])
        if not spec:
            events.append({'filename':source['filename'],'source_date':source['as_of'],'status':'unsupported','rows':len(source['rows']),'matched':0,'updated_cells':0});continue
        for book,target_name in spec['targets']:
            event={'filename':source['filename'],'sha256':source['sha256'],'source_date':source['as_of'],'book':book,'target_sheet':target_name,'rows':len(source['rows']),'matched':0,'updated_cells':0,'pending':0,'status':'checked'}
            events.append(event);target=sheets.get((book,target_name))
            if not source['as_of']:event['status']='date_unknown';continue
            if source['as_of']>as_of:event['status']='future_date';continue
            if source['as_of']<baseline_by_book[book]:event['status']='stale';continue
            if not target:event['status']='target_missing';continue
            tc,td=column_map(target['columns'])
            for duplicate in td:
                candidates=[c['col'] for c in target['columns'] if header(c['label'])==duplicate]
                raw=[col for col in candidates if not any(c.get('formula') for r in target['rows'] if r['row']>target['header_row'] for c in r['cells'] if c['col']==col)]
                if len(raw)==1:tc[duplicate]=raw[0]
            if any(header(k) not in sc or header(k) not in tc for k in spec['keys']):event['status']='key_header_missing';continue
            # Headers missing from one side are never mapped by physical position.
            event['unmapped_headers']=[c['label'] for c in source['columns'] if header(c['label']) not in tc]
            event['ambiguous_headers']=sorted(set(sd+td))
            ti=defaultdict(list);si=Counter(identity(r,sc,spec['keys']) for r in source['rows'])
            for r in target['rows']:
                if r['row']>target['header_row']:ti[identity(r,tc,spec['keys'])].append(r)
            for row in source['rows']:
                key=identity(row,sc,spec['keys']);matches=ti.get(key,[]) if key else []
                if not key or si[key]!=1 or len(matches)!=1:
                    reason='missing_identity' if not key else 'duplicate_identity' if si[key]>1 or len(matches)>1 else 'new_or_unmatched'
                    pending.append({'file':source['filename'],'book':book,'sheet':target_name,'source_row':row['row'],'reason':reason,'identity':list(key or [])});event['pending']+=1;continue
                target_row=matches[0];before={c['col']:c for c in target_row['cells']};incoming={c['col']:c for c in row['cells']};event['matched']+=1
                stale=False
                for date_label in ['처리일자','처리일']:
                    label=header(date_label)
                    if label in sc and label in tc:
                        old_date=normalized(before.get(tc[label],{}).get('value'));new_date=normalized(incoming.get(sc[label],{}).get('value'))
                        if re.match(r'^\d{4}-\d{2}-\d{2}',old_date) and re.match(r'^\d{4}-\d{2}-\d{2}',new_date) and new_date[:10]<old_date[:10]:stale=True
                if stale:
                    pending.append({'file':source['filename'],'book':book,'sheet':target_name,'source_row':row['row'],'reason':'older_record_date','identity':list(key)});event['pending']+=1;continue
                for label,source_col in sc.items():
                    if label not in tc:continue
                    target_col=tc[label];old=before.get(target_col,{});new=incoming.get(source_col,{}).get('value')
                    if old.get('formula'):continue
                    if config.get('auto_update_headers') and label not in {header(h) for h in config['auto_update_headers']}:
                        if normalized(new)!=normalized(old.get('value')):
                            pending.append({'file':source['filename'],'book':book,'sheet':target_name,'source_row':row['row'],'target_row':target_row['row'],'column':target_col,'reason':'non_status_change_review','identity':list(key)});event['pending']+=1
                        continue
                    if new is None or new=='':
                        pending.append({'file':source['filename'],'book':book,'sheet':target_name,'source_row':row['row'],'target_row':target_row['row'],'column':target_col,'reason':'blank_does_not_clear','identity':list(key)});event['pending']+=1;continue
                    if isinstance(new,str) and new.lstrip().startswith(('=','+','@')):
                        pending.append({'file':source['filename'],'book':book,'sheet':target_name,'source_row':row['row'],'reason':'formula_like_value','identity':list(key)});event['pending']+=1;continue
                    w={'book':book,'sheet':target_name,'row':target_row['row'],'col':target_col,'before':old.get('value'),'value':new,'source_file':source['filename'],'source_row':row['row'],'source_date':source['as_of'],'header':label}
                    claims[(book,target_name,target_row['row'],target_col)].append((w,event))
    for key,items in claims.items():
        # Even equal-date conflicts must not silently use the last file.
        latest=max(w['source_date'] for w,e in items);eligible=[(w,e) for w,e in items if w['source_date']==latest]
        if len({normalized(w['value']) for w,e in eligible})!=1:
            for w,e in eligible:pending.append({'file':w['source_file'],'book':w['book'],'sheet':w['sheet'],'source_row':w['source_row'],'reason':'conflicting_exports','identity':[]});e['pending']+=1
            continue
        w,e=eligible[0];writes.append(w);e['updated_cells']+=1
        if normalized(w['value'])==normalized(w['before']):writes.pop();e['updated_cells']-=1
    return {'schema_version':'ors-refresh-1.0','writes':writes,'files':events,'pending':pending,
        'summary':{'source_files':len({s['sha256'] for s in exports}),'file_target_checks':len(events),'updated_cells':len(writes),'pending_rows':len(pending),'stale_file_targets':sum(e['status']=='stale' for e in events)}}
