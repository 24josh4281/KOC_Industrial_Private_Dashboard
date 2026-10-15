"""Cross-check the v14 display files against the unchanged source catalog."""
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from ors_model import collect, packed_json, summarize

ROOT=Path(__file__).resolve().parent
V13=next(p for p in (ROOT.parent/'v13',ROOT.parent/'2026-10-15_meeting_v13') if p.exists())
folder=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT
source=V13 if folder==ROOT else V13/'snapshots'/folder.name
data=json.loads((source/'dashboard-data.json').read_text(encoding='utf-8'))
catalog=packed_json(source/'ors-full-catalog.gz.b64')
rows=collect(data,catalog);summary=summarize(data,rows)
checks={'source_occurrences':len(catalog),'ors_rows':len(rows),'unique_ids':len({x['id'] for x in rows}),
        'latest_rows':summary['latest_ors_rows'],'industry_latest':summary['latest_industry_rows'],
        'industry_2026_intake':summary['ors_2026_intake_rows_industry'],
        'industry_2026_activity':summary['ors_2026_activity_rows_industry']}
assert checks['ors_rows']==sum(len(item['rows'])-1 for item in catalog)
assert checks['unique_ids']==checks['ors_rows']
assert all(x['source_row']>=2 and x['raw_excel_sheet'] for x in rows)
for name,expected in [('통계_대시보드_v14.html',summary),('진행현황_대시보드_v14.html',summary)]:
    page=(folder/name).read_text(encoding='utf-8')
    assert '__DATA__' not in page and '__CSS__' not in page
    assert '외부사업_통합관리_v14.xlsx' in page
    scripts=re.findall(r'<script(?:\s[^>]*)?>(.*?)</script>',page,re.S)
    assert len(scripts)==2,(name,len(scripts))
    payload=json.loads(scripts[0])
    assert payload['summary']['source_occurrence_rows']==summary['source_occurrence_rows']
    if 'rows' in payload:
        assert len(payload['rows'])==len(rows)
        assert len(payload['management'])==len(data['cases_2026'])
        assert len({x['id'] for x in payload['rows']})==len(rows)
    temp=ROOT/'work'/('check_'+name+'.js');temp.parent.mkdir(exist_ok=True)
    temp.write_text(scripts[1],encoding='utf-8')
    node=Path('C:/Users/24jos/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
    if node.exists():subprocess.run([str(node),'--check',str(temp)],check=True,capture_output=True)
    checks[name]={'bytes':(folder/name).stat().st_size,'javascript_syntax':'ok' if node.exists() else 'runtime unavailable'}
book=load_workbook(folder/'외부사업_통합관리_v14.xlsx',read_only=True,data_only=False)
assert book.sheetnames[:2]==['통계 대시보드','진행현황 대시보드']
assert book['ORS 전체 진행'].max_row==len(rows)+1
assert book['ORS 2026 인입'].max_row==sum(1 for x in rows if x['is_latest'] and x['intake_dates_2026'])+1
checks['excel_sheets']=len(book.sheetnames)
checks['excel_all_ors_rows']=book['ORS 전체 진행'].max_row-1
checks['excel_2026_intake_rows']=book['ORS 2026 인입'].max_row-1
book.close()
with zipfile.ZipFile(folder/'외부사업_통합관리_v14.xlsx') as archive:
    assert archive.testzip() is None
    checks['excel_charts']=sum('/charts/chart' in n and n.endswith('.xml') for n in archive.namelist())
assert checks['excel_charts']>=9
if folder==ROOT:
    assert checks['source_occurrences']==20 and checks['ors_rows']==5998
    assert checks['latest_rows']==4197 and checks['industry_latest']==2895
    assert checks['industry_2026_intake']==27 and checks['industry_2026_activity']==169
print(json.dumps(checks,ensure_ascii=False,indent=2))
