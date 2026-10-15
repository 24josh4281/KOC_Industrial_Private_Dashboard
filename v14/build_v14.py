"""Build separate statistics and progress dashboards plus a matching Excel.

Reads the checked v13 snapshot and ORS source catalog without modifying either.
Monthly use: --db PATH first creates/reuses a v13 DB snapshot, then adds v14.
"""
import argparse
import json
import os
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from ors_model import collect, packed_json, summarize

ROOT = Path(__file__).resolve().parent
V13 = next((p for p in (ROOT.parent/'v13', ROOT.parent/'2026-10-15_meeting_v13') if p.exists()), None)
NAVY, TEAL, BLUE, PALE, WHITE = '102A43', '087F8C', '3559D9', 'EAF1F5', 'FFFFFF'
KIND_NAMES = {
    'domestic_project':'국내 사업신청','domestic_credit':'국내 인증신청',
    'foreign_project':'국외 사업신청','foreign_credit':'국외 인증신청',
    'disclosure_project_domestic':'비공개 사업승인','disclosure_credit_domestic':'비공개 인증승인',
    'disclosure_project_foreign':'국외 비공개 사업','disclosure_credit_foreign':'국외 비공개 인증',
    'unclassified':'분류 확인'}


def safe(value):
    if isinstance(value,str):
        if value[:1] in ('=','+','@'):
            value="'"+value
        return value[:32767]
    return value


def scope_stats(progress, scope):
    latest=[x for x in progress if x['is_latest'] and (scope=='전체' or x['scope']==scope)]
    intake=[x for x in latest if x['intake_dates_2026']]
    activity=[x for x in latest if x['event_count_2026']]
    month=Counter(date[5:7] for x in intake for date in x['intake_dates_2026'])
    status=Counter(x['status_bucket'] for x in latest if x['intake_dates_2026'] or x['event_count_2026'])
    return {'latest_rows':len(latest),'intake_rows_2026':len(intake),
            'intake_events_2026':sum(len(x['intake_dates_2026']) for x in intake),
            'activity_rows_2026':len(activity),
            'intake_month_2026':{f'{i:02d}':month.get(f'{i:02d}',0) for i in range(1,13)},
            'status_2026':dict(status.most_common()),
            'by_kind':dict(Counter(KIND_NAMES.get(x['kind'],x['kind']) for x in latest))}


def management_cases(data):
    output=[]
    for case in data['cases_2026']:
        note=data['workflow'].get(case['case_key'],{})
        output.append({**{k:case.get(k) for k in ('case_key','title','source_file','source_row','apply_date','raw_status')},
                       **{k:note.get(k,'') for k in ('intake_status','opinion_status','opinion_text',
                                                   'guidance_status','guidance_text','next_action')}})
    return output


def source_inputs(args):
    if V13 is None:
        raise FileNotFoundError('v13 기준 자료 폴더를 찾을 수 없습니다. 전체 저장소를 내려받으세요.')
    if args.db:
        p=subprocess.run([sys.executable,str(V13/'refresh_outputs.py'),'--db',str(args.db)],
                         capture_output=True,text=True,encoding='utf-8',check=True)
        last=json.loads(p.stdout.strip().splitlines()[-1])
        folder=Path(last['snapshot'])
        base=folder/'외부사업_통합관리_전체자료.xlsx'
        out=ROOT/'snapshots'/folder.name
    else:
        folder=V13
        base=folder/'외부사업_통합관리_전체자료_v13_검증완료.xlsx'
        out=ROOT
    for path in (folder/'dashboard-data.json',folder/'ors-full-catalog.gz.b64',base):
        if not path.exists():
            raise FileNotFoundError(path)
    return folder,base,out


def html(template, css, payload, raw_link):
    content=template.read_text(encoding='utf-8')
    content=(content.replace('__ORS_ROWS__',f"{payload['summary']['source_occurrence_rows']:,}")
             .replace('__MANAGEMENT_ROWS__',f"{payload['summary']['management_2026_rows']:,}")
             .replace('__ORS_SHEETS__',f"ORS01~{payload['ors_file_occurrences']:02d}"))
    content=content.replace('__CSS__',css).replace('__DATA__',json.dumps(payload,ensure_ascii=False,separators=(',',':')).replace('<','\\u003c'))
    return content.replace('../v13/대시보드_v13.html',raw_link)


def section_head(ws,title,subtitle):
    ws.sheet_view.showGridLines=False
    ws.merge_cells('A1:P2');ws['A1']=title
    ws['A1'].fill=PatternFill('solid',fgColor=NAVY)
    ws['A1'].font=Font(name='Malgun Gothic',size=21,bold=True,color=WHITE)
    ws['A1'].alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=23;ws.row_dimensions[2].height=24
    ws.merge_cells('A3:P3');ws['A3']=subtitle;ws['A3'].font=Font(name='Malgun Gothic',size=10,color='50667A')
    ws.row_dimensions[3].height=26
    for col in 'ABCDEFGHIJKLMNOP':ws.column_dimensions[col].width=13
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1


def card(ws,range_text,label,value,note,color=NAVY):
    left=range_text.split(':')[0]
    row=int(''.join(c for c in left if c.isdigit()))
    col=''.join(c for c in left if c.isalpha())
    end=range_text.split(':')[-1]
    last_col=''.join(c for c in end if c.isalpha())
    ws.merge_cells(f'{col}{row}:{last_col}{row}')
    ws.merge_cells(f'{col}{row+1}:{last_col}{row+2}')
    ws.merge_cells(f'{col}{row+3}:{last_col}{row+3}')
    for cells in ws[f'{col}{row}:{last_col}{row+3}']:
        for c in cells:c.fill=PatternFill('solid',fgColor=color)
    ws[f'{col}{row}']=label;ws[f'{col}{row}'].font=Font(name='Malgun Gothic',size=10,color='DDECF5',bold=True)
    ws[f'{col}{row+1}']=value;ws[f'{col}{row+1}'].font=Font(name='Malgun Gothic',size=24,color=WHITE,bold=True)
    ws[f'{col}{row+3}']=note;ws[f'{col}{row+3}'].font=Font(name='Malgun Gothic',size=9,color='DDECF5')
    for r in range(row,row+4):ws.row_dimensions[r].height=max(ws.row_dimensions[r].height or 0,22)


def title_row(ws,row,title):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16)
    c=ws.cell(row,1,title);c.fill=PatternFill('solid',fgColor=PALE)
    c.font=Font(name='Malgun Gothic',size=12,bold=True,color=NAVY)
    c.alignment=Alignment(vertical='center')
    ws.row_dimensions[row].height=27


def chart(ws,source,cat_col,val_col,min_row,max_row,anchor,title,color):
    graph=BarChart();graph.title=title;graph.style=10;graph.height=8;graph.width=16
    graph.y_axis.title='원본행';graph.x_axis.title=None
    graph.add_data(Reference(source,min_col=val_col,min_row=min_row,max_row=max_row),titles_from_data=True)
    graph.set_categories(Reference(source,min_col=cat_col,min_row=min_row+1,max_row=max_row))
    graph.legend=None
    if graph.series:graph.series[0].graphicalProperties.solidFill=color
    ws.add_chart(graph,anchor)


def data_table(ws,headers,rows,name,widths=None):
    ws.sheet_view.showGridLines=False;ws.freeze_panes='D2';ws.append(headers)
    for c in ws[1]:c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Malgun Gothic',bold=True,color=WHITE,size=10);c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=34
    for row in rows:ws.append([safe(x) for x in row])
    if rows:
        ref=f'A1:{ws.cell(1,len(headers)).column_letter}{len(rows)+1}'
        table=Table(displayName=name,ref=ref)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    for i,w in enumerate(widths or [18]*len(headers),1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
    ws.auto_filter.ref=f'A1:{ws.cell(1,len(headers)).column_letter}{len(rows)+1}'


def excel(base,out,data,progress,summary,scopes):
    wb=load_workbook(base)
    for name in ('통계 대시보드','진행현황 대시보드','ORS 전체 진행','ORS 2026 인입','집계 근거'):
        if name in wb.sheetnames:raise ValueError(f'이미 존재하는 v14 시트: {name}')
    stats=wb.create_sheet('통계 대시보드',0);active=wb.create_sheet('진행현황 대시보드',1)
    all_rows=wb.create_sheet('ORS 전체 진행');intake_ws=wb.create_sheet('ORS 2026 인입');evidence=wb.create_sheet('집계 근거')
    latest=[x for x in progress if x['is_latest']]
    industry=scopes['산업']
    data_table(all_rows,['최신본','수신일','범위','유형','사업·신청명','신청자','번호','신청유형','원본 처리상태','상태 묶음','2026 접수일','접수 근거','2026 이력수','최근 이력일','원본 파일','원본 시트','원본 행','Excel 원본 시트'],
               [[('Y' if x['is_latest'] else 'N'),x['snapshot_date'],x['scope'],KIND_NAMES.get(x['kind'],x['kind']),x['title'],x['applicant'],x['number'],x['application_type'],x['status'],x['status_bucket'],' / '.join(x['intake_dates_2026']),' / '.join(x['intake_basis_2026']),x['event_count_2026'],x['last_event_date'],x['source_file'],x['source_sheet'],x['source_row'],x['raw_excel_sheet']] for x in progress],
               'OrsProgressAll',[10,14,12,20,58,28,22,16,33,17,18,27,13,14,55,32,12,17])
    intake_rows=[x for x in latest if x['intake_dates_2026']]
    data_table(intake_ws,['범위','유형','2026 접수일','접수 근거','사업·신청명','신청자','현재 표시 상태','2026 처리 이력수','최근 이력','원본 파일','원본 행','Excel 원본 시트'],
               [[x['scope'],KIND_NAMES.get(x['kind'],x['kind']),' / '.join(x['intake_dates_2026']),' / '.join(x['intake_basis_2026']),x['title'],x['applicant'],x['status'],x['event_count_2026'],x['last_event'],x['source_file'],x['source_row'],x['raw_excel_sheet']] for x in intake_rows],
               'OrsIntake2026',[12,20,18,26,58,28,35,16,58,55,12,17])
    for ws,rows,source_col in ((all_rows,progress,18),(intake_ws,intake_rows,12)):
        for i,x in enumerate(rows,2):
            cell=ws.cell(i,source_col)
            cell.hyperlink=f"#'{x['raw_excel_sheet']}'!A{x['source_row']+4}"
            cell.font=Font(name='Malgun Gothic',color=BLUE,underline='single')
    data_table(evidence,['연도','관리 Tool 원장행','월','관리 Tool 2026','ORS 산업 2026 접수','상태','ORS 산업 행','유형','ORS 산업 최신 행'],
               [[str(y),summary['management_yearly'][str(y)],f'{i:02d}',summary['management_month_2026'][f'{i:02d}'],industry['intake_month_2026'][f'{i:02d}'],
                 (list(industry['status_2026'])+[None]*12)[i-1],(list(industry['status_2026'].values())+[None]*12)[i-1],
                 (list(industry['by_kind'])+[None]*12)[i-1],(list(industry['by_kind'].values())+[None]*12)[i-1]]
                for i,y in enumerate(range(2015,2027),1)],'Evidence2026',[12,18,12,18,20,24,18,22,20])
    section_head(stats,'통계 대시보드 | 누적·2026 인입',
                 '관리 Tool 원장행과 ORS 최신 수신행을 구분. 서로 합산해 고유 신청 건수로 해석하지 않습니다.')
    card(stats,'A5:C8','관리 Tool 누적 원장행',summary['management_rows'],'연도 미확인·2014 이전 포함')
    card(stats,'E5:G8','관리 Tool 2026 신청일',summary['management_2026_rows'],'신청일 확인 · 원장행',BLUE)
    card(stats,'I5:K8','ORS 산업 최신 원본행',industry['latest_rows'],'유형별 최신 파일',TEAL)
    card(stats,'M5:O8','ORS 산업 2026 접수',industry['intake_rows_2026'],'접수 근거 있는 원본행')
    title_row(stats,10,'01  관리 Tool 연도별 추이');title_row(stats,29,'02  2026 월별 관리 Tool / ORS 산업 비교')
    chart(stats,evidence,1,2,1,13,'A11','연도별 신청 원장행',BLUE)
    chart(stats,evidence,3,4,1,13,'A30','관리 Tool 월별 신청일 행',BLUE)
    chart(stats,evidence,3,5,1,13,'I30','ORS 산업 월별 접수 근거행',TEAL)
    stats.merge_cells('I11:P15');stats['I11']=(
        '주의: ORS 2026 접수 근거는 명시 신청일 또는 신규·갱신 접수/최초 신청 이력입니다. '
        f"보완·재제출은 제외하며, 관리 Tool {summary['management_2026_rows']:,}행과 "
        '신청 회차별 자동 연결은 미확정입니다.')
    stats['I11'].alignment=Alignment(wrap_text=True,vertical='center');stats['I11'].fill=PatternFill('solid',fgColor=PALE)
    stats['I11'].font=Font(name='Malgun Gothic',size=11,color=NAVY)
    stats['A49']='상세 숫자: 집계 근거 시트 · 원본 셀: 원본01~36 및 ORS01~';stats['A49'].font=Font(name='Malgun Gothic',size=10,color=NAVY)
    section_head(active,'진행현황 대시보드 | ORS·2026 업무',
                 f"ORS 원본행 전체 {summary['source_occurrence_rows']:,}개 출현 보존. 기본 KPI는 유형·범위별 최신 산업부문 파일 기준입니다.")
    card(active,'A5:C8','ORS 전체 수신 출현',summary['source_occurrence_rows'],'월별 원본 ZIP 누적')
    card(active,'E5:G8','ORS 산업 최신',industry['latest_rows'],'과거 스냅샷 중복 제외',BLUE)
    card(active,'I5:K8','산업 2026 접수 근거',industry['intake_rows_2026'],'고유 신청 수 아님',TEAL)
    card(active,'M5:O8','관리 Tool 2026 업무',summary['management_2026_rows'],'의견·안내 별도 관리')
    title_row(active,10,'01  산업부문 2026 처리 이력 행의 현재 상태');title_row(active,29,'02  산업부문 2026 접수 월별 분포')
    chart(active,evidence,6,7,1,13,'A11','현재 표시 상태',TEAL)
    chart(active,evidence,3,5,1,13,'A30','2026 접수 근거 월별 행',BLUE)
    active.merge_cells('I11:P16');active['I11']=(
        f"ORS 산업 2026 이력·직접 날짜가 있는 {summary['ors_2026_activity_rows_industry']:,}행에는 "
        '과거 접수 후 보완·재제출·승인도 포함됩니다. ‘ORS 2026 인입’에는 접수 근거가 있는 '
        f"최신 원본행만 표시합니다. 관리 Tool {summary['management_2026_rows']:,}행은 별도 원장입니다.")
    active['I11'].alignment=Alignment(wrap_text=True,vertical='center');active['I11'].fill=PatternFill('solid',fgColor=PALE);active['I11'].font=Font(name='Malgun Gothic',size=11,color=NAVY)
    active['I30']='전체 상세: ORS 전체 진행 / 올해 접수: ORS 2026 인입 / 내부 의견: 2026 접수·검토·안내';active['I30'].alignment=Alignment(wrap_text=True);active['I30'].font=Font(name='Malgun Gothic',size=10,color=NAVY)
    for ws in (stats,active):ws.freeze_panes='A4';ws.sheet_view.zoomScale=85
    out.parent.mkdir(parents=True,exist_ok=True)
    temp=out.with_suffix('.building.xlsx')
    try:
        wb.save(temp)
        with zipfile.ZipFile(temp) as archive:
            bad=archive.testzip();charts=sum('/charts/chart' in n and n.endswith('.xml') for n in archive.namelist())
        if bad or charts<8:raise RuntimeError(f'Excel 차트/ZIP 검사 실패: bad={bad}, charts={charts}')
        os.replace(temp,out)
    finally:
        wb.close()
    return {'sheets':len(wb.sheetnames),'charts':charts,'ors_rows':len(progress),'ors_2026_intake_rows':len(intake_rows)}


def main():
    parser=argparse.ArgumentParser(description='통계·진행현황 분리 v14 산출')
    parser.add_argument('--db',type=Path,help='운영 DB; 먼저 v13 동기 판본을 생성')
    parser.add_argument('--data-only',action='store_true',help='집계 검토용 JSON만 생성')
    args=parser.parse_args()
    folder,base,out=source_inputs(args)
    data=json.loads((folder/'dashboard-data.json').read_text(encoding='utf-8'))
    catalog=packed_json(folder/'ors-full-catalog.gz.b64')
    progress=collect(data,catalog);summary=summarize(data,progress)
    scopes={scope:scope_stats(progress,scope) for scope in ('산업','기후','발전','전체')}
    as_of=data.get('refresh_snapshot',{}).get('latest_ors_date') or data['as_of']
    management=management_cases(data)
    output={'as_of':as_of,'summary':summary,'scopes':scopes,
            'management_count':len(management),'ors_file_occurrences':len(catalog)}
    out.mkdir(parents=True,exist_ok=True)
    (out/'검증수치_v14.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    if args.data_only:
        print(json.dumps({'output':str(out),'summary':{k:v for k,v in summary.items() if not isinstance(v,dict)}},ensure_ascii=False));return
    css=(ROOT/'dashboard.css').read_text(encoding='utf-8')
    raw_target=folder/'대시보드.html' if args.db else V13/'대시보드_v13.html'
    raw_link=os.path.relpath(raw_target,out).replace('\\','/')
    (out/'통계_대시보드_v14.html').write_text(html(ROOT/'stats_template.html',css,
           {'as_of':as_of,'summary':summary,'scopes':scopes,
            'ors_file_occurrences':len(catalog)},raw_link),encoding='utf-8')
    (out/'진행현황_대시보드_v14.html').write_text(html(ROOT/'progress_template.html',css,
           {'as_of':as_of,'summary':summary,'rows':progress,'management':management,
            'ors_file_occurrences':len(catalog)},raw_link),encoding='utf-8')
    xlsx=out/'외부사업_통합관리_v14.xlsx'
    check=excel(base,xlsx,data,progress,summary,scopes)
    print(json.dumps({'output':str(out),'summary':{k:v for k,v in summary.items() if not isinstance(v,dict)},
                      'excel':check},ensure_ascii=False))


if __name__=='__main__':main()
