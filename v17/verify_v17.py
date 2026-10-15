"""Check that the v17 redesign changes presentation without losing source data."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


CHANGED_STYLES = {'통계 대시보드', '진행현황 대시보드', '업무 운영 대시보드',
                  '담당자 작업현황', '요구사항 점검'}
EXPECTED_CHART_COLORS = {
    '통계 대시보드': ('2F62F0', '2F62F0', '244FC0'),
    '진행현황 대시보드': ('244FC0', '2F62F0'),
    '업무 운영 대시보드': ('244FC0', '2F62F0', '2F62F0'),
}


def contrast(foreground, background):
    def luminance(value):
        channels = [int(value[i:i + 2], 16) / 255 for i in (1, 3, 5)]
        linear = [c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
                  for c in channels]
        return sum(a * b for a, b in zip(linear, (0.2126, 0.7152, 0.0722)))

    a, b = sorted((luminance(foreground), luminance(background)), reverse=True)
    return (a + 0.05) / (b + 0.05)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def chart_refs(chart):
    refs = []
    for series in chart.series:
        value = series.val.numRef.f if series.val and series.val.numRef else None
        category = None
        if series.cat:
            if series.cat.numRef:
                category = series.cat.numRef.f
            elif series.cat.strRef:
                category = series.cat.strRef.f
        refs.append((value, category))
    return refs


def verify(site: Path):
    old = site / 'v16'
    new = site / 'v17'
    prior = json.loads((old / 'manifest.json').read_text(encoding='utf-8'))
    manifest = json.loads((new / 'manifest.json').read_text(encoding='utf-8'))
    for key in ('as_of', 'ors_row_occurrences', 'management_2026', 'work_source_rows'):
        if prior[key] != manifest[key]:
            raise AssertionError(f'숫자·기준일 변경: {key}')
    if manifest['v16_version'] != prior['version'] or manifest['brand_profile'] != 'innergen_blue':
        raise AssertionError('디자인 판본 출처가 다릅니다.')
    for group in ('data', 'details'):
        before = sorted((old / group).glob('*.json'))
        after = sorted((new / group).glob('*.json'))
        if [x.name for x in before] != [x.name for x in after]:
            raise AssertionError(f'{group} 파일 목록 변경')
        for left, right in zip(before, after, strict=True):
            if sha(left) != sha(right):
                raise AssertionError(f'원본 데이터 변경: {right}')
    css = (new / 'brand.css').read_text(encoding='utf-8')
    for token in ('#0B1739', '#2F62F0', '#7B9BF7', '#EEF2FF', '#E2E8F0'):
        if token not in css:
            raise AssertionError(f'이너젠 디자인 색 누락: {token}')
    brand = json.loads((Path(__file__).with_name('brand.json')).read_text(encoding='utf-8'))
    for role in ('navy', 'blue', 'dark_blue', 'ink', 'muted'):
        if contrast(brand[role], '#FFFFFF') < 4.5:
            raise AssertionError(f'흰 바탕 텍스트 대비 부족: {role}')
    for name in ('statistics.html', 'progress.html', 'work.html', 'case.html'):
        html = (new / name).read_text(encoding='utf-8')
        if ('href="brand.css"' not in html or
                '외부사업_통합관리_v16.xlsx' in html or
                '외부사업_통합관리_v15.xlsx' in html):
            raise AssertionError(f'웹 스타일 또는 Excel 연결 누락: {name}')
        if name != 'case.html' and '<span class="stamp-label">자동 갱신 ·</span>' not in html:
            raise AssertionError(f'갱신 기준일 표기 누락: {name}')

    a = load_workbook(old / '외부사업_통합관리_v16.xlsx')
    b = load_workbook(new / '외부사업_통합관리_v17.xlsx')
    try:
        if a.sheetnames != b.sheetnames or len(b.sheetnames) != 77:
            raise AssertionError('시트 구조 변경')
        compared = 0
        unchanged_style_cells = 0
        for old_ws, new_ws in zip(a.worksheets, b.worksheets, strict=True):
            if old_ws.max_row != new_ws.max_row or old_ws.max_column != new_ws.max_column:
                raise AssertionError(f'셀 범위 변경: {old_ws.title}')
            if (old_ws.freeze_panes != new_ws.freeze_panes or
                    old_ws.auto_filter.ref != new_ws.auto_filter.ref or
                    set(old_ws.tables) != set(new_ws.tables)):
                raise AssertionError(f'표·탐색 구조 변경: {old_ws.title}')
            for old_row, new_row in zip(old_ws, new_ws, strict=True):
                for left, right in zip(old_row, new_row, strict=True):
                    compared += 1
                    if left.value != right.value:
                        # Six shortened labels on the presentation-only work sheet.
                        if old_ws.title == '업무 운영 대시보드' and left.coordinate in {
                                'A11', 'D11', 'G11', 'J11', 'A27', 'A39'}:
                            continue
                        raise AssertionError(f'값·수식 변경: {old_ws.title}!{left.coordinate}')
                    if old_ws.title not in CHANGED_STYLES:
                        unchanged_style_cells += 1
                        if left._style != right._style:
                            raise AssertionError(f'원본 스타일 변경: {old_ws.title}!{left.coordinate}')
            if len(old_ws._charts) != len(new_ws._charts):
                raise AssertionError(f'차트 수 변경: {old_ws.title}')
            for before, after in zip(old_ws._charts, new_ws._charts, strict=True):
                if chart_refs(before) != chart_refs(after):
                    raise AssertionError(f'차트 연결 변경: {old_ws.title}')
        chart_count = sum(len(ws._charts) for ws in b.worksheets)
        if chart_count != 12:
            raise AssertionError(f'차트 12개 유지 실패: {chart_count}')
        for sheet_name, colors in EXPECTED_CHART_COLORS.items():
            ws = b[sheet_name]
            if ws['A1'].fill.fgColor.rgb[-6:] != '0B1739':
                raise AssertionError(f'Excel 남색 제목 누락: {sheet_name}')
            for chart, color in zip(ws._charts, colors, strict=True):
                if chart.series[0].graphicalProperties.solidFill.srgbClr != color:
                    raise AssertionError(f'차트 색 변경 실패: {sheet_name}')
        if b['진행현황 대시보드']._charts[0].type != 'bar':
            raise AssertionError('긴 상태 레이블 개선 실패')
        return {'sheets': len(b.sheetnames), 'charts': chart_count,
                'compared_cells': compared, 'unchanged_style_cells': unchanged_style_cells,
                'work_rows': manifest['work_source_rows'], 'profile': manifest['brand_profile']}
    finally:
        a.close()
        b.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--site', required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(verify(args.site), ensure_ascii=False))
